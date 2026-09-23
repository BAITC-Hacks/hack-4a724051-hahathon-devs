import io
import zipfile
from pathlib import Path

import pytest

from app.modules.documents.extraction import MAX_TABLE_ROWS, extract
from app.modules.documents.validation import FileRejected, check_file, safe_filename

CERT_PDF = sorted((Path(__file__).resolve().parents[2] / "data" / "synthetic" / "certificates").glob("*.pdf"))[0]


def xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx_bytes(paragraphs, table=None):
    import docx
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def zip_bytes(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


SHEET_TYPES = '<Types><Override ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/></Types>'


def test_detects_kind_by_content_not_name():
    assert check_file("spec.pdf", xlsx([["a"]])).kind == "xlsx"
    assert check_file("x.xlsx", docx_bytes(["hi"])).kind == "docx"
    assert check_file("x", CERT_PDF.read_bytes()).kind == "pdf"
    assert check_file("x", b"\xff\xd8\xff\xe0" + b"0" * 10).kind == "jpeg"


@pytest.mark.parametrize("data,code", [
    (b"", "empty_file"),
    (b"MZ\x90\x00 not a document", "unsupported_type"),
    (b"\xd0\xcf\x11\xe0 old xls", "legacy_office"),
    (b"%PDF-" + b"0" * (11 * 1024 * 1024), "file_too_large"),
], ids=["empty", "executable", "legacy-office", "oversize"])
def test_rejects_bad_files(data, code):
    with pytest.raises(FileRejected) as e:
        check_file("f", data)
    assert e.value.code == code


def test_rejects_macros_and_zip_bomb():
    with pytest.raises(FileRejected, match="macros"):
        check_file("m.xlsm", zip_bytes({"[Content_Types].xml": SHEET_TYPES, "xl/vbaProject.bin": b"x"}))
    with pytest.raises(FileRejected) as e:
        check_file("b.xlsx", zip_bytes({"[Content_Types].xml": SHEET_TYPES, "xl/big.xml": b"0" * 70_000_000}))
    assert e.value.code == "too_large_unpacked"


def test_safe_filename():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("C:\\docs\\спец<>.xlsx") == "спец_.xlsx"


def test_xlsx_rows_and_limit():
    small = extract("xlsx", xlsx([["Артикул", "Кол-во"], ["990100001_", 10], [None, None]]))
    assert small.status == "ready"
    assert "990100001_ | 10" in small.text
    big = extract("xlsx", xlsx([[f"row{i}", i] for i in range(MAX_TABLE_ROWS + 50)]))
    assert big.status == "partial"
    assert "ещё 50 не обработано" in big.warnings[0]


def test_docx_paragraphs_and_tables():
    result = extract("docx", docx_bytes(["Спецификация щита"], [["Автомат 16А", "10 шт"]]))
    assert result.status == "ready"
    assert "Спецификация щита" in result.text and "Автомат 16А | 10 шт" in result.text


def test_pdf_text_layer():
    result = extract("pdf", CERT_PDF.read_bytes())
    assert result.status == "ready"
    assert "SYNTHETIC CERTIFICATE" in result.text


def test_images_need_ocr():
    result = extract("jpeg", b"\xff\xd8\xff")
    assert (result.status, result.error) == ("failed", "needs_ocr")


def test_broken_file_fails_without_crashing():
    broken = zip_bytes({"[Content_Types].xml": SHEET_TYPES})
    assert check_file("b.xlsx", broken).kind == "xlsx"
    result = extract("xlsx", broken)
    assert (result.status, result.error) == ("failed", "parse_error")
