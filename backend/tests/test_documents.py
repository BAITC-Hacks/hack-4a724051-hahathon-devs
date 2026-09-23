import asyncio
import io
import sqlite3
import sys
import time
import zipfile
from pathlib import Path

import pytest

from app.core.errors import AppError
from app.modules.documents import DocumentsService


def service(tmp_path: Path, scanner: bool) -> DocumentsService:
    db_path = tmp_path / "state.sqlite3"
    with sqlite3.connect(db_path) as db:
        db.execute("CREATE TABLE sessions (id TEXT PRIMARY KEY, token_hash TEXT, expires_at REAL NOT NULL)")
        db.execute("INSERT INTO sessions(id,expires_at) VALUES (?,?)", ("owner", time.time() + 3600))
        db.execute("INSERT INTO sessions(id,expires_at) VALUES (?,?)", ("other", time.time() + 3600))
    command = None
    if scanner:
        script = tmp_path / "scan.py"
        script.write_text("import sys\nsys.exit(0)\n", encoding="utf-8")
        command = [sys.executable, str(script)]
    result = DocumentsService(db_path, tmp_path / "assets", command)
    result.initialize()
    return result


def run(coro):
    return asyncio.run(coro)


def test_scanner_env_keeps_windows_paths_but_never_secrets(monkeypatch):
    monkeypatch.setenv("SYSTEMROOT", "C:/Windows")
    monkeypatch.setenv("ALLUSERSPROFILE", "C:/ProgramData")
    monkeypatch.setenv("LLM_API_KEY", "must-not-inherit")
    environment = DocumentsService._subprocess_env()
    assert environment["SYSTEMROOT"] == "C:/Windows"
    assert environment["ALLUSERSPROFILE"] == "C:/ProgramData"
    assert "LLM_API_KEY" not in environment


@pytest.mark.parametrize("code", [1, 2, 50])
def test_scanner_nonzero_never_releases_file(tmp_path, monkeypatch, code):
    import subprocess
    documents = service(tmp_path, True)
    monkeypatch.setattr(subprocess, "run", lambda *a, **kw: subprocess.CompletedProcess(a[0], code))
    asset = run(documents.upload("owner", "parts.txt", "text/plain", b"990100003_"))
    assert asset.status == "quarantined"
    with pytest.raises(AppError):
        run(documents.resolve("owner", [str(asset.id)]))


def test_scanner_disabled_quarantines_and_cannot_resolve(tmp_path):
    documents = service(tmp_path, False)
    asset = run(documents.upload("owner", "parts.txt", "text/plain", b"ABC-123"))
    assert asset.status == "quarantined"
    assert not (tmp_path / "assets" / "clean" / str(asset.id)).exists()
    with pytest.raises(AppError) as error:
        run(documents.resolve("owner", [str(asset.id)]))
    assert error.value.code == "document_quarantined"
    with pytest.raises(AppError) as error:
        run(documents.get("other", str(asset.id)))
    assert error.value.status == 404
    run(documents.delete("owner", str(asset.id)))
    assert not (tmp_path / "assets" / "quarantine" / str(asset.id)).exists()


def test_clean_text_owned_order_and_duplicates(tmp_path):
    documents = service(tmp_path, True)
    one = run(documents.upload("owner", "one.txt", "text/plain", b"First article"))
    two = run(documents.upload("owner", "two.csv", "text/csv", b"article,qty\nA1,2\n"))
    assert one.status == two.status == "ready"
    resolved = run(documents.resolve("owner", [str(two.id), str(one.id), str(two.id)]))
    assert [item.asset_id for item in resolved] == [two.id, one.id, two.id]
    assert "A1" in resolved[0].text
    with pytest.raises(AppError) as error:
        run(documents.resolve("other", [str(one.id)]))
    assert error.value.status == 404


def test_signature_and_ooxml_active_content_rejected(tmp_path):
    documents = service(tmp_path, True)
    with pytest.raises(AppError) as error:
        run(documents.upload("owner", "bad.pdf", "application/pdf", b"not a pdf"))
    assert error.value.status == 415
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(AppError) as error:
        run(documents.upload("owner", "bad.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", output.getvalue()))
    assert error.value.code == "unsafe_document"


def test_real_parsers_and_image_budget(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    from pypdf import PdfWriter

    documents = service(tmp_path, True)
    workbook = Workbook()
    workbook.active.append(["Артикул", "Количество"])
    workbook.active.append(["A-42", 3])
    spreadsheet = io.BytesIO()
    workbook.save(spreadsheet)
    xlsx = run(documents.upload("owner", "parts.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", spreadsheet.getvalue()))
    assert "A-42" in run(documents.resolve("owner", [str(xlsx.id)]))[0].text

    word = Document()
    word.add_paragraph("Кабель ВВГнг")
    docx_bytes = io.BytesIO()
    word.save(docx_bytes)
    docx = run(documents.upload("owner", "request.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_bytes.getvalue()))
    assert "Кабель" in run(documents.resolve("owner", [str(docx.id)]))[0].text

    image = Image.new("RGB", (2100, 100), "white")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")
    png = run(documents.upload("owner", "scan.png", "image/png", image_bytes.getvalue()))
    owned = run(documents.images("owner", [str(xlsx.id), str(png.id)]))
    assert len(owned) == 1 and owned[0].asset_id == png.id
    with Image.open(io.BytesIO(owned[0].content)) as normalized:
        assert max(normalized.size) == 2000

    pdf = PdfWriter()
    pdf.add_blank_page(width=300, height=300)
    pdf_bytes = io.BytesIO()
    pdf.write(pdf_bytes)
    scanned = run(documents.upload("owner", "scan.pdf", "application/pdf", pdf_bytes.getvalue()))
    parsed = run(documents.resolve("owner", [str(scanned.id)]))[0]
    assert parsed.status == "partial"
    assert "изображени" in " ".join(parsed.warnings)
    assert run(documents.images("owner", [str(scanned.id)]))[0].mime_type == "image/png"


def test_quota_reservation_and_expired_cleanup(tmp_path, monkeypatch):
    import app.modules.documents.service as module
    documents = service(tmp_path, False)
    monkeypatch.setattr(module, "MAX_SESSION_ASSETS", 1)
    first = run(documents.upload("owner", "first.txt", "text/plain", b"one"))
    with pytest.raises(AppError) as error:
        run(documents.upload("owner", "second.txt", "text/plain", b"two"))
    assert error.value.code == "asset_quota"
    with sqlite3.connect(documents.db_path) as db:
        db.execute("UPDATE sessions SET expires_at=? WHERE id='owner'", (time.time() - 1,))
    assert run(documents.cleanup()) == 1
    assert not (documents.root_path / "quarantine" / str(first.id)).exists()


def test_payment_text_deleted_after_rejection(tmp_path):
    documents = service(tmp_path, True)
    with pytest.raises(AppError) as error:
        run(documents.upload("owner", "payment.txt", "text/plain", b"card 4111 1111 1111 1111"))
    assert error.value.code == "sensitive_data_rejected"
    with sqlite3.connect(documents.db_path) as db:
        assert db.execute("SELECT count(*) FROM document_assets").fetchone()[0] == 0
    assert not list((documents.root_path / "quarantine").iterdir())
    assert not list((documents.root_path / "clean").iterdir())
