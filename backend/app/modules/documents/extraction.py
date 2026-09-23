"""Извлечение текста из xlsx, docx и pdf.

Парсер работает в отдельном процессе с таймаутом и лимитом памяти: зависший или
упавший разбор одного файла не трогает сервер. Лимиты страниц, строк и символов
не скрываются: если что-то не прочитано, статус partial и понятное предупреждение.
Картинки и сканы без текстового слоя требуют OCR, которого пока нет.
"""

import io
import multiprocessing as mp
from dataclasses import dataclass, field

MAX_TEXT_CHARS = 40000
MAX_PDF_PAGES = 20
MAX_TABLE_ROWS = 200
PARSE_TIMEOUT_S = 30
MEMORY_LIMIT_BYTES = 1024 * 1024 * 1024


@dataclass
class Extraction:
    status: str  # ready | partial | failed
    text: str = ""
    warnings: list[str] = field(default_factory=list)
    error: str | None = None  # код для клиента, если failed


def _clip(parts: list[str], warnings: list[str]) -> str:
    text = "\n".join(p for p in parts if p).strip()
    if len(text) > MAX_TEXT_CHARS:
        warnings.append(f"Текст обрезан до {MAX_TEXT_CHARS} символов.")
        text = text[:MAX_TEXT_CHARS]
    return text


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\n", " ").strip()


def parse_xlsx(data: bytes) -> Extraction:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, warnings, rows_used, skipped = [], [], 0, 0
    for sheet in wb.worksheets:
        parts.append(f"# Лист: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            if not any(cells):
                continue
            if rows_used >= MAX_TABLE_ROWS:
                skipped += 1
                continue
            while cells and not cells[-1]:
                cells.pop()
            parts.append(" | ".join(cells))
            rows_used += 1
    wb.close()
    if skipped:
        warnings.append(f"Прочитано {MAX_TABLE_ROWS} строк, ещё {skipped} не обработано.")
    text = _clip(parts, warnings)
    if not rows_used:
        return Extraction("failed", error="no_text", warnings=["В таблице нет данных."])
    return Extraction("partial" if warnings else "ready", text, warnings)


def parse_docx(data: bytes) -> Extraction:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts, warnings = [], []
    parts.extend(p.text.strip() for p in document.paragraphs if p.text.strip())
    rows = [row for table in document.tables for row in table.rows]
    for row in rows[:MAX_TABLE_ROWS]:
        cells = [c.text.replace("\n", " ").strip() for c in row.cells]
        if any(cells):
            parts.append(" | ".join(cells))
    if len(rows) > MAX_TABLE_ROWS:
        warnings.append(f"Из таблиц прочитано {MAX_TABLE_ROWS} строк из {len(rows)}.")
    text = _clip(parts, warnings)
    if not text:
        return Extraction("failed", error="no_text", warnings=["В документе нет текста."])
    return Extraction("partial" if warnings else "ready", text, warnings)


def parse_pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        return Extraction("failed", error="encrypted", warnings=["PDF защищён паролем."])
    total = len(reader.pages)
    parts, warnings, empty_pages = [], [], 0
    for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"# Страница {number}\n{text}")
        else:
            empty_pages += 1
    if total > MAX_PDF_PAGES:
        warnings.append(f"Прочитано {MAX_PDF_PAGES} страниц из {total}.")
    if empty_pages:
        warnings.append(f"Страниц без текстового слоя (скан): {empty_pages}. Их содержимое не прочитано.")
    text = _clip(parts, warnings)
    if not text:
        return Extraction("failed", error="needs_ocr",
                          warnings=["В PDF нет текстового слоя, похоже на скан. Распознавание сканов пока не подключено."])
    return Extraction("partial" if warnings else "ready", text, warnings)


PARSERS = {"xlsx": parse_xlsx, "docx": parse_docx, "pdf": parse_pdf}


def _child(kind: str, data: bytes, queue) -> None:
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (MEMORY_LIMIT_BYTES, MEMORY_LIMIT_BYTES))
    except (ImportError, ValueError, OSError):
        pass  # не на всех ОС лимит памяти доступен
    try:
        queue.put(PARSERS[kind](data))
    except Exception as e:  # noqa: BLE001 - любой сбой парсера это failed, а не падение сервера
        queue.put(Extraction("failed", error="parse_error", warnings=[f"Не удалось разобрать файл ({type(e).__name__})."]))


def extract(kind: str, data: bytes, timeout_s: float = PARSE_TIMEOUT_S) -> Extraction:
    if kind in ("jpeg", "png"):
        return Extraction("failed", error="needs_ocr",
                          warnings=["Распознавание текста на фото пока не подключено. Пришлите артикулы текстом "
                                    "или файлом Excel/Word/PDF."])
    if kind not in PARSERS:
        return Extraction("failed", error="unsupported_type", warnings=["Формат не поддерживается."])
    ctx = mp.get_context("spawn")
    queue = ctx.Queue(maxsize=1)
    process = ctx.Process(target=_child, args=(kind, data, queue), daemon=True)
    process.start()
    try:
        result = queue.get(timeout=timeout_s)
    except Exception:  # noqa: BLE001 - queue.Empty или обрыв дочернего процесса
        result = Extraction("failed", error="parse_timeout", warnings=["Разбор файла занял слишком много времени."])
    finally:
        process.join(1)
        if process.is_alive():
            process.kill()
            process.join()
    return result
